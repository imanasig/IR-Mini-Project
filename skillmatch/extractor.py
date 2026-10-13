"""Read supported documents in memory; never execute or persist uploads."""
from io import BytesIO
from pathlib import Path
import fitz
from docx import Document

MAX_BYTES = 5 * 1024 * 1024

class ExtractionError(ValueError):
    pass

def extract_text(name, content):
    if len(content) > MAX_BYTES:
        raise ExtractionError('File exceeds the 5 MB limit.')
    suffix = Path(name).suffix.lower()
    if suffix not in {'.txt', '.pdf', '.docx'}:
        raise ExtractionError('Use a PDF, DOCX, or TXT document.')
    try:
        if suffix == '.txt':
            text = content.decode('utf-8-sig')
        elif suffix == '.pdf':
            with fitz.open(stream=content, filetype='pdf') as doc:
                text = '\n'.join(page.get_text() for page in doc)
        else:
            doc = Document(BytesIO(content))
            blocks = [p.text for p in doc.paragraphs]
            blocks.extend(' | '.join(c.text for c in row.cells)
                          for table in doc.tables for row in table.rows)
            text = '\n'.join(blocks)
    except Exception as exc:
        raise ExtractionError('Unable to read this file. It may be corrupted, encrypted, or not UTF-8 text.') from exc
    if not text.strip():
        message = 'No readable text found.'
        if suffix == '.pdf':
            message += ' Scanned PDFs require OCR, which this version does not support.'
        raise ExtractionError(message)
    return text.strip()
